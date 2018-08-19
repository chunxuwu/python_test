# -*- coding: utf-8 -*-
"""
Created on Sun Aug 19 15:36:10 2018
12.10.4 图表
AttributeError: module 'openpyxl' has no attribute 'charts'
解决方法：
https://blog.csdn.net/weixin_41569319/article/details/80821715

程序是能运行，只是，画的图加载不出来
@author: NEVERGUVEIP
"""
import openpyxl
wb = openpyxl.Workbook()
sheet = wb.active
for i in range(1,11):#creat some data in column A
    sheet['A'+str(i)] = i

refObj = openpyxl.chart.Reference(sheet,min_row = 1,min_col = 1,max_row =10,max_col = 1)

seriesObj = openpyxl.chart.Series(refObj,title='First series')

chartObj = openpyxl.chart.BarChart()

chartObj.title = 'My Chart'

chartObj.add_data(refObj)
sheet.add_chart(chartObj,'C5')
#chartObj.drawing.top = 50
#chartObj.drawing.keft = 100
#chartObj.drawing.width = 300
#chartObj.drawing.height = 200

sheet.add_chart(chartObj)
wb.save('sampleChart.xlsx')
